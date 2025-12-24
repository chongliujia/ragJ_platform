"""
Parser Service
This service is responsible for parsing text content from various file types.
Integrates with Rust document processor for improved performance.
"""

import fitz  # PyMuPDF
import docx
from io import BytesIO
import logging
from typing import Optional, Dict, Any
from bs4 import BeautifulSoup
import markdown as md

logger = logging.getLogger(__name__)

# Try to import Rust document processor
try:
    from app.services.rust_document_service import rust_processor

    RUST_AVAILABLE = rust_processor is not None
    if RUST_AVAILABLE:
        logger.info(
            "Rust document processor available - using high-performance parsing"
        )
    else:
        logger.info(
            "Rust document processor not available - falling back to Python parsing"
        )
except ImportError:
    RUST_AVAILABLE = False
    logger.info("Rust document processor not installed - using Python parsing")


def parse_document(
    content: bytes, filename: str, options: Optional[Dict[str, Any]] = None
) -> str:
    """
    Parse document using Rust processor if available, otherwise fallback to Python.

    Args:
        content: Document content as bytes
        filename: Original filename (used for format detection)
        options: Parsing options (OCR, table extraction, etc.)

    Returns:
        Extracted text content
    """
    if RUST_AVAILABLE:
        try:
            return rust_processor.parse_content(content, filename, options)
        except Exception as e:
            logger.warning(
                f"Rust parsing failed for {filename}, falling back to Python: {e}"
            )
            return _parse_document_python(content, filename)
    else:
        return _parse_document_python(content, filename)


def _parse_document_python(content: bytes, filename: str) -> str:
    """Fallback Python-based document parsing."""
    file_ext = filename.lower().split(".")[-1] if "." in filename else ""

    if file_ext == "txt":
        return parse_txt(content)
    elif file_ext == "pdf":
        return parse_pdf(content)
    elif file_ext == "docx":
        return parse_docx(content)
    elif file_ext == "md":
        return parse_md(content)
    elif file_ext in {"xlsx", "xls"}:
        return parse_excel(content, filename)
    elif file_ext == "html" or file_ext == "htm":
        return parse_html(content)
    else:
        logger.warning(f"Unsupported file format: {file_ext}")
        return ""


def parse_txt(content: bytes) -> str:
    """Parses text from a .txt file."""
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        logger.warning("Failed to decode TXT file with UTF-8, trying with gbk.")
        try:
            return content.decode("gbk")
        except UnicodeDecodeError:
            logger.error("Failed to decode TXT file with both UTF-8 and gbk.")
            return ""


def parse_pdf(content: bytes) -> str:
    """Parses text from a .pdf file."""
    try:
        doc = fitz.open(stream=content, filetype="pdf")
        text = ""
        for page in doc:
            text += page.get_text()
        return text
    except Exception as e:
        logger.error(f"Error parsing PDF file: {e}", exc_info=True)
        return ""


def parse_docx(content: bytes) -> str:
    """Parses text from a .docx file."""
    try:
        document = docx.Document(BytesIO(content))
        text = "\n".join([para.text for para in document.paragraphs])
        return text
    except Exception as e:
        logger.error(f"Error parsing DOCX file: {e}", exc_info=True)
        return ""


def parse_md(content: bytes) -> str:
    """Parses text from a .md (Markdown) file by converting to HTML then stripping tags."""
    try:
        text = content.decode("utf-8", errors="ignore")
        # Convert markdown to HTML
        html = md.markdown(text)
        # Strip HTML tags to plain text
        soup = BeautifulSoup(html, "html.parser")
        return soup.get_text(separator="\n")
    except Exception as e:
        logger.error(f"Error parsing MD file: {e}", exc_info=True)
        return ""


def parse_html(content: bytes) -> str:
    """Parses text from a .html file using BeautifulSoup."""
    try:
        html = content.decode("utf-8", errors="ignore")
        soup = BeautifulSoup(html, "html.parser")
        # Remove script/style elements
        for tag in soup(["script", "style"]):
            tag.decompose()
        return soup.get_text(separator="\n")
    except Exception as e:
        logger.error(f"Error parsing HTML file: {e}", exc_info=True)
        return ""


def parse_excel(content: bytes, filename: str) -> str:
    """Parses text from an Excel file (.xlsx/.xls)."""
    file_ext = filename.lower().split(".")[-1] if "." in filename else ""
    if file_ext == "xlsx":
        try:
            import openpyxl

            wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
            lines = []
            for sheet in wb.worksheets:
                lines.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = []
                    for cell in row:
                        if cell is None:
                            continue
                        if isinstance(cell, float) and cell.is_integer():
                            values.append(str(int(cell)))
                        else:
                            values.append(str(cell))
                    if values:
                        lines.append("\t".join(values))
            return "\n".join(lines)
        except Exception as e:
            logger.error(f"Error parsing XLSX file: {e}", exc_info=True)
            return ""
    if file_ext == "xls":
        try:
            import xlrd

            wb = xlrd.open_workbook(file_contents=content)
            lines = []
            for sheet in wb.sheets():
                lines.append(f"Sheet: {sheet.name}")
                for row_idx in range(sheet.nrows):
                    values = []
                    for cell in sheet.row_values(row_idx):
                        if cell in ("", None):
                            continue
                        if isinstance(cell, float) and cell.is_integer():
                            values.append(str(int(cell)))
                        else:
                            values.append(str(cell))
                    if values:
                        lines.append("\t".join(values))
            return "\n".join(lines)
        except Exception as e:
            logger.error(f"Error parsing XLS file: {e}", exc_info=True)
            return ""
    logger.warning(f"Unsupported Excel file format: {file_ext}")
    return ""


def get_supported_formats() -> list[str]:
    """Get list of supported document formats."""
    if RUST_AVAILABLE:
        try:
            return rust_processor.supported_formats
        except Exception as e:
            logger.warning(f"Failed to get supported formats from Rust: {e}")

    # Fallback to Python-supported formats
    return ["txt", "pdf", "docx", "md", "html", "htm", "xlsx", "xls"]


def extract_metadata(content: bytes, filename: str) -> Dict[str, Any]:
    """Extract metadata from document."""
    if RUST_AVAILABLE:
        try:
            return rust_processor.get_metadata_from_content(content, filename)
        except Exception as e:
            logger.warning(f"Rust metadata extraction failed: {e}")

    # Basic metadata fallback
    return {
        "file_size": len(content),
        "filename": filename,
        "format": filename.lower().split(".")[-1] if "." in filename else "unknown",
    }
